# -*- coding: utf-8 -*-
"""Convierte docs/estudio/unidad-1-guia-extensa.md a Word para estudiar."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "docs" / "estudio" / "unidad-1-guia-extensa.md"
OUT = ROOT / "docs" / "estudio" / "unidad-1-guia-extensa.docx"

NAVY = RGBColor(0x10, 0x26, 0x3D)
AMBER = RGBColor(0xE5, 0x9A, 0x24)
MUTED = RGBColor(0x55, 0x55, 0x55)
FONT = "Calibri"


def set_run_font(run, size=11, bold=False, italic=False, color=None, name=FONT):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:cs"), name)
    rFonts.set(qn("w:eastAsia"), name)


def shade(cell, hex_color: str) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag == qn("w:shd"):
            tcPr.remove(child)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def add_inline(paragraph, text: str, size=11, color=None) -> None:
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size=size, bold=True, color=color or NAVY)
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, size=size - 1, italic=False, color=MUTED, name="Consolas")
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, size=size, italic=True, color=color)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=size, color=color)


def heading(doc, text: str, level: int) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level <= 2 else 10)
    p.paragraph_format.space_after = Pt(6)
    sizes = {1: 20, 2: 16, 3: 13, 4: 12}
    colors = {1: NAVY, 2: NAVY, 3: AMBER, 4: NAVY}
    run = p.add_run(text)
    set_run_font(run, size=sizes.get(level, 12), bold=True, color=colors.get(level, NAVY))
    if level == 1:
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "12")
        bottom.set(qn("w:space"), "4")
        bottom.set(qn("w:color"), "E59A24")
        pBdr.append(bottom)
        pPr.append(pBdr)


def para(doc, text: str, *, quote=False, bullet=False, number=None) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    if quote:
        p.paragraph_format.left_indent = Cm(0.75)
        add_inline(p, text, size=11, color=MUTED)
        if p.runs:
            p.runs[0].italic = True
        return
    if bullet:
        p.paragraph_format.left_indent = Cm(0.75)
        add_inline(p, "•  " + text)
        return
    if number is not None:
        p.paragraph_format.left_indent = Cm(0.75)
        add_inline(p, f"{number}.  " + text)
        return
    add_inline(p, text)


def code_block(doc, lines: list[str]) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Cm(0.4)
    run = p.add_run("\n".join(lines))
    set_run_font(run, size=8.5, name="Consolas", color=NAVY)


def add_table(doc, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(r) for r in rows)
    for r in rows:
        while len(r) < cols:
            r.append("")
    tbl = doc.add_table(rows=len(rows), cols=cols)
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, cell_text in enumerate(row):
            cell = tbl.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            add_inline(p, cell_text.strip(), size=9, color=RGBColor(0xFF, 0xFF, 0xFF) if i == 0 else None)
            if i == 0:
                shade(cell, "10263D")
                for run in p.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            elif i % 2 == 1:
                shade(cell, "F7F3EA")
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(8)


def is_sep_row(line: str) -> bool:
    core = line.strip().strip("|")
    return bool(core) and all(set(part.strip()) <= set("-: ") for part in core.split("|"))


def parse_table_row(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def convert(md: str) -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.2)
    header = section.header.paragraphs[0]
    hr = header.add_run("Planificación de proyectos de software · Guía de estudio")
    set_run_font(hr, size=9, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fr = footer.add_run("Definiciones para estudio · lenguaje PMBOK 6.ª")
    set_run_font(fr, size=8, color=MUTED)

    lines = md.splitlines()
    i = 0
    in_code = False
    code_acc: list[str] = []

    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()

        if line.startswith("```"):
            if in_code:
                code_block(doc, code_acc)
                code_acc = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_acc.append(raw)
            i += 1
            continue

        if not line.strip() or line.strip() == "---":
            i += 1
            continue

        if line.startswith("# "):
            heading(doc, line[2:].strip(), 1)
            i += 1
            continue
        if line.startswith("## "):
            heading(doc, line[3:].strip(), 2)
            i += 1
            continue
        if line.startswith("### "):
            heading(doc, line[4:].strip(), 3)
            i += 1
            continue
        if line.startswith("#### "):
            heading(doc, line[5:].strip(), 4)
            i += 1
            continue

        if line.lstrip().startswith("|") and i + 1 < len(lines) and is_sep_row(lines[i + 1]):
            rows = [parse_table_row(line)]
            i += 2
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                if not is_sep_row(lines[i]):
                    rows.append(parse_table_row(lines[i]))
                i += 1
            add_table(doc, rows)
            continue

        if line.startswith("> "):
            quote_parts = [line[2:].strip()]
            i += 1
            while i < len(lines) and lines[i].startswith("> "):
                quote_parts.append(lines[i][2:].strip())
                i += 1
            para(doc, " ".join(quote_parts), quote=True)
            continue

        m_num = re.match(r"^(\d+)\.\s+(.*)$", line)
        if m_num:
            para(doc, m_num.group(2), number=m_num.group(1))
            i += 1
            continue

        if line.startswith("- "):
            para(doc, line[2:].strip(), bullet=True)
            i += 1
            continue

        para(doc, line.strip())
        i += 1

    return doc


def main() -> None:
    md = SRC.read_text(encoding="utf-8")
    doc = convert(md)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
