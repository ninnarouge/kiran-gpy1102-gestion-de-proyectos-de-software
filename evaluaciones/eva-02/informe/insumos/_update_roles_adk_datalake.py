# -*- coding: utf-8 -*-
"""Nicolás y Ari: GCloud, ADK (agéntico) y datalake como base de información."""
from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

DOCX = Path(__file__).resolve().parents[1] / "2.2.2-Gestion-de-cambios-Kiran v2.docx"


def set_cell_text(cell, text: str) -> None:
    if not cell.paragraphs:
        cell.text = text
        return
    p0 = cell.paragraphs[0]
    if p0.runs:
        p0.runs[0].text = text
        for run in p0.runs[1:]:
            run.text = ""
    else:
        p0.text = text
    for p in cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)


def set_para(p: Paragraph, text: str) -> None:
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.text = text


def find_para(doc: Document, contains: str = None, startswith: str = None):
    for p in doc.paragraphs:
        t = p.text or ""
        if startswith is not None and t.startswith(startswith):
            return p
        if contains is not None and contains in t:
            return p
    raise KeyError(contains or startswith)


def insert_paragraph_after(paragraph: Paragraph, text: str) -> Paragraph:
    new_p = deepcopy(paragraph._p)
    for child in list(new_p):
        if child.tag == qn("w:r"):
            new_p.remove(child)
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if paragraph.runs:
        run = new_para.add_run(text)
        src = paragraph.runs[0]
        run.bold = src.bold
        run.italic = src.italic
        if src.font.size:
            run.font.size = src.font.size
        if src.font.name:
            run.font.name = src.font.name
    else:
        new_para.add_run(text)
    return new_para


def replace_contains(doc: Document, needle: str, old: str, new: str) -> bool:
    for p in doc.paragraphs:
        t = p.text or ""
        if needle in t and old in t:
            set_para(p, t.replace(old, new))
            return True
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                t = cell.text or ""
                if needle in t and old in t:
                    set_cell_text(cell, t.replace(old, new))
                    return True
    return False


def main() -> None:
    doc = Document(str(DOCX))

    # --- §2.5 cuerpo C-05 ---
    for p in doc.paragraphs:
        t = p.text or ""
        if not t.startswith("En los términos pedidos"):
            continue
        if "chatbot" not in t.lower() and "FAQ" not in t:
            continue
        set_para(
            p,
            "En los términos pedidos (chatbot para usuarios con FAQ y comunicación de "
            "soporte, para ahorrar contrato de personal), el cambio toca la interfaz de la "
            "plataforma, la claridad de la experiencia de uso, la integración con Gemini en "
            "Google Cloud, la base de información (datalake) que alimenta FAQ, monitoreos "
            "futuros y reportes, la privacidad, la calidad de las respuestas y el flujo hacia "
            "el ticket. Giannina sostiene la interfaz del chat y la UI/UX asociada. Skarlett "
            "valida la calidad de las respuestas y el FAQ. Nicolás Barra y Ari Araya quedan "
            "encargados de lo técnico de datos e inteligencia de esta solución: Google Cloud, "
            "Agent Development Kit (ADK) para lo agéntico, y el datalake como base de "
            "información del piloto (ingesta acotada, almacenamiento y lectura para el chat y "
            "para fases posteriores de monitoreo). Un lakehouse enterprise completo, una app "
            "para cada hogar o el monitoreo total de flota con IA no se construyen de golpe en "
            "esta ventana: se diseñan sobre esa base a cargo de Nicolás y Ari, con control de "
            "cambios y financiamiento explícito.",
        )
        break

    # --- Hoja de ruta / responsabilidad ---
    inserted = False
    for p in doc.paragraphs:
        if "Responsabilidad técnica agéntica" in (p.text or "") or (
            "Agent Development Kit" in (p.text or "") and "Nicolás Barra y Ari" in (p.text or "")
        ):
            set_para(
                p,
                "Responsabilidad técnica (datos e inteligencia): Nicolás Barra y Ari Araya "
                "quedan a cargo de Google Cloud, del Agent Development Kit (ADK) para el "
                "desarrollo agéntico, y del datalake como base de información del piloto "
                "(inventario, tickets, logs de chat y trazas que después alimenten monitoreo y "
                "alertas). No se contrata un perfil externo para eso en sep–dic 2026. "
                "Giannina mantiene el control de cambios y la UI; Skarlett, la calidad del FAQ "
                "y de las respuestas. El datalake se trata como fundamento: en esta ventana "
                "se prioriza un depósito mínimo usable (sin PII innecesaria) alineado al "
                "chatbot; la ampliación a lakehouse o analítica pesada sigue sujeta a CC y "
                "techo (Escenario C u otro).",
            )
            inserted = True
            break

    if not inserted:
        try:
            anchor = find_para(doc, contains="Hoja de ruta del hilo C-05")
        except KeyError:
            try:
                anchor = find_para(doc, contains="efecto positivo acotado")
            except KeyError:
                anchor = find_para(doc, startswith="2.5. C-05")
        insert_paragraph_after(
            anchor,
            "Responsabilidad técnica (datos e inteligencia): Nicolás Barra y Ari Araya "
            "quedan a cargo de Google Cloud, del Agent Development Kit (ADK) para el "
            "desarrollo agéntico, y del datalake como base de información del piloto "
            "(inventario, tickets, logs de chat y trazas para monitoreo futuro). No se "
            "contrata un perfil externo para eso en esta ventana. Giannina: control de "
            "cambios y UI. Skarlett: calidad del FAQ. El datalake se prioriza como "
            "fundamento mínimo usable en la ventana; el lakehouse completo sigue con CC.",
        )

    # --- §3.1 C-05 ---
    try:
        p = find_para(doc, startswith="C-05 se aprueba solo por fases")
        t = p.text or ""
        if "datalake" not in t.lower() or "ADK" not in t:
            extra = (
                " En lo técnico de datos e inteligencia (Google Cloud, ADK y datalake como "
                "base de información), los responsables son Nicolás Barra y Ari Araya; no se "
                "incorpora especialista externo en esta ventana."
            )
            if "Nicolás Barra y Ari Araya" not in t or "datalake" not in t.lower():
                set_para(p, t.rstrip() + extra)
    except KeyError:
        pass

    # --- Efecto conjunto: datalake importante, no solo rechazo ---
    try:
        p = find_para(doc, contains="datalake o lakehouse")
        t = p.text or ""
        if "base de información" not in t:
            set_para(
                p,
                t.rstrip()
                + " El datalake, no obstante, se reconoce como base de información importante "
                "para el chatbot, el soporte y el monitoreo futuro: queda a cargo de Nicolás "
                "y Ari. En esta ventana se apunta a un depósito mínimo usable; el lakehouse "
                "completo no se aprueba sin CC y financiamiento.",
            )
    except KeyError:
        pass

    # Soften blanket "datalake rechazado" where it says completo only
    replace_contains(
        doc,
        "datalake",
        "Un datalake completo, un especialista agéntico externo",
        "Un lakehouse enterprise sin CC, un especialista agéntico externo",
    )

    # --- Tabla 8 ---
    for table in doc.tables:
        if (table.rows[0].cells[0].text or "") != "Paso":
            continue
        for row in table.rows:
            p0 = row.cells[0].text or ""
            if "C-05 fase 1" in p0 or "chatbot FAQ" in p0:
                set_cell_text(
                    row.cells[1],
                    "Giannina Guerrero (CC e interfaz del chat); Nicolás Barra y Ari Araya "
                    "(Google Cloud, ADK, datalake como base de información e integración "
                    "Gemini); Skarlett Tropan (calidad FAQ); Ari Araya (gasto y topes Excel)",
                )
            if "ADK" in p0 or "Agent Development Kit" in p0:
                set_cell_text(
                    row.cells[0],
                    "Definir y mantener en Google Cloud el entorno del chatbot, los "
                    "componentes agénticos con Agent Development Kit (ADK) cuando el alcance "
                    "lo exija, y el datalake como base de información (cuotas, secretos, "
                    "logs, límites de gasto Gemini y datos mínimos sin PII innecesaria).",
                )
                set_cell_text(
                    row.cells[1],
                    "Nicolás Barra y Ari Araya (co-responsables: GCloud, ADK y datalake)",
                )
        # Si no hay fila ADK/datalake, crearla
        if not any(
            "datalake" in (r.cells[0].text or "").lower() or "ADK" in (r.cells[0].text or "")
            for r in table.rows
        ):
            tr = deepcopy(table.rows[-1]._tr)
            table._tbl.append(tr)
            set_cell_text(
                table.rows[-1].cells[0],
                "Definir y mantener Google Cloud, ADK (lo agéntico) y el datalake como base "
                "de información del piloto: cuotas, secretos, logs, límites Gemini y depósito "
                "mínimo usable sin PII innecesaria.",
            )
            set_cell_text(
                table.rows[-1].cells[1],
                "Nicolás Barra y Ari Araya (co-responsables técnicos)",
            )
            set_cell_text(
                table.rows[-1].cells[2],
                "Al activar C-05b (tras financiamiento) y en cada revisión quincenal de gasto "
                "y de datos mientras el chat esté en uso.",
            )
        break

    # --- Tabla 10 C-05a / C-05b ---
    for table in doc.tables:
        if len(table.columns) < 6:
            continue
        if (table.rows[0].cells[0].text or "") != "Id.":
            continue
        for row in table.rows:
            id_ = row.cells[0].text or ""
            if id_.startswith("C-05b"):
                set_cell_text(
                    row.cells[2],
                    "Nicolás Barra y Ari Araya (GCloud, ADK, datalake); Giannina (CC)",
                )
                just = row.cells[5].text or ""
                if "datalake" not in just.lower():
                    set_cell_text(
                        row.cells[5],
                        just.rstrip()
                        + " GCloud, ADK y datalake (base de información) a cargo de Nicolás y "
                        "Ari; sin especialista externo. Datalake mínimo alineado al chatbot en "
                        "esta ventana; ampliación con CC.",
                    )
            if id_.startswith("C-05a"):
                # Reencuadrar: no solo rechazo total del datalake
                tipo = row.cells[1].text or ""
                if "base de información" not in tipo:
                    set_cell_text(
                        row.cells[1],
                        "Requisito futuro / ampliación. Chatbot amplio; app ciudadana; "
                        "reemplazo total del soporte; monitoreo y alertas con IA; lakehouse "
                        "completo. El datalake como base de información queda identificado y "
                        "a cargo de Nicolás y Ari (mínimo en C-05b; ampliación aquí).",
                    )
                set_cell_text(
                    row.cells[2],
                    "Nicolás Barra y Ari Araya (GCloud, ADK, datalake); Giannina (CC)",
                )
                just = row.cells[5].text or ""
                if "Nicolás" not in just or "datalake" not in just.lower():
                    set_cell_text(
                        row.cells[5],
                        "No se aprueba la formulación amplia en sep–dic 2026 sin CC y techo. "
                        "Dueños técnicos si se reabre: Nicolás y Ari (GCloud, ADK, datalake). "
                        "El datalake mínimo de C-05b no se confunde con este lakehouse completo.",
                    )
        # ¿Existe fila datalake? Si no, agregar C-05c o ampliar - user wants datalake important
        # Add C-07 or note in C-05 - better add row C-05d datalake base
        ids = [r.cells[0].text for r in table.rows]
        if not any("datalake" in (i or "").lower() or i.startswith("C-05d") for i in ids):
            tr = deepcopy(table.rows[-1]._tr)
            table._tbl.append(tr)
            set_cell_text(table.rows[-1].cells[0], "C-05d")
            set_cell_text(
                table.rows[-1].cells[1],
                "Datos. Datalake como base de información del piloto (mínimo usable: "
                "inventario, tickets, logs de chat agregados; sin PII innecesaria).",
            )
            set_cell_text(
                table.rows[-1].cells[2],
                "Nicolás Barra y Ari Araya",
            )
            set_cell_text(table.rows[-1].cells[3], "Alta")
            set_cell_text(table.rows[-1].cells[4], "Abierto")
            set_cell_text(
                table.rows[-1].cells[5],
                "Aprobado en alcance mínimo alineado a C-05b (base para FAQ y fases "
                "posteriores). Lakehouse completo no. Financiamiento dentro del Escenario C "
                "o CC si el storage supera el margen GCloud. Dueños: Nicolás y Ari.",
            )
        break

    # --- Tabla 2: si hay C-05, mencionar datalake dueños en texto bajo tabla ---
    try:
        p = find_para(doc, startswith="C-05 se rotula")
        t = p.text or ""
        if "datalake" not in t.lower():
            set_para(
                p,
                t.rstrip()
                + " El datalake como base de información del piloto queda a cargo de Nicolás "
                "y Ari, junto con Google Cloud y ADK.",
            )
    except KeyError:
        pass

    # --- Tabla 7d CAPEX dueños ---
    for table in doc.tables:
        blob = " ".join(c.text for r in table.rows for c in r.cells)
        if "CAPEX" not in blob and "puesta en marcha" not in blob:
            continue
        for row in table.rows:
            c0 = row.cells[0].text or ""
            if "Integración Gemini" in c0 or "despliegue" in c0.lower():
                set_cell_text(
                    row.cells[0],
                    "Integración Gemini, ADK, datalake mínimo y despliegue en Google Cloud",
                )
                if len(row.cells) > 4:
                    set_cell_text(
                        row.cells[4],
                        "Nicolás Barra y Ari Araya (GCloud, ADK, datalake)",
                    )
        break

    # --- Tabla 7b: fila storage datalake si cabe ---
    for table in doc.tables:
        blob = " ".join(c.text for r in table.rows for c in r.cells)
        if "Subtotal caja C-05b" not in blob:
            continue
        # Ampliar fila GCloud para mencionar datalake mínimo
        for row in table.rows:
            if "Google Cloud" in (row.cells[0].text or ""):
                set_cell_text(
                    row.cells[0],
                    "Google Cloud (chat, logs, datalake mínimo, cuotas)",
                )
                set_cell_text(
                    row.cells[4],
                    "Módulo chat + depósito mínimo de información (Nicolás/Ari). Sin "
                    "lakehouse enterprise. Sin GPU.",
                )
        break

    # --- Mensaje cliente ---
    try:
        p = find_para(doc, contains="Mensaje al cliente")
        t = p.text or ""
        if "datalake" not in t.lower():
            set_para(
                p,
                t.rstrip()
                + " El datalake como base de información, el ADK y Google Cloud los ejecutan "
                "Nicolás y Ari dentro del equipo.",
            )
    except KeyError:
        pass

    # --- Tabla 11 ---
    t11 = doc.tables[-1]
    if "Pregunta" in (t11.rows[0].cells[0].text or ""):
        for row in t11.rows:
            q = row.cells[0].text or ""
            if "CAPEX" in q:
                a = row.cells[1].text or ""
                if "datalake" not in a.lower():
                    set_cell_text(
                        row.cells[1],
                        a.rstrip()
                        + " Datalake mínimo, ADK y GCloud: Nicolás y Ari.",
                    )
        if not any("quién" in (r.cells[0].text or "").lower() and "datalake" in (r.cells[0].text or "").lower() for r in t11.rows):
            if not any("ADK" in (r.cells[0].text or "") for r in t11.rows):
                tr = deepcopy(t11.rows[-1]._tr)
                t11._tbl.append(tr)
                set_cell_text(
                    t11.rows[-1].cells[0],
                    "¿Quién sostiene GCloud, ADK y el datalake?",
                )
                set_cell_text(
                    t11.rows[-1].cells[1],
                    "Nicolás Barra y Ari Araya, co-responsables técnicos. El datalake se "
                    "trata como base de información del piloto (mínimo en esta ventana; "
                    "ampliación con CC). Sin especialista externo en sep–dic 2026.",
                )

    # Intro
    try:
        p = find_para(doc, contains="C-05 y C-06 no suman encima")
        t = p.text or ""
        if "datalake como base" not in t:
            set_para(
                p,
                t.rstrip()
                + " El datalake se reconoce como base de información a cargo de Nicolás y Ari "
                "(mínimo usable con C-05b; lakehouse completo solo con CC).",
            )
    except KeyError:
        pass

    doc.save(str(DOCX))
    print("OK", DOCX)


if __name__ == "__main__":
    main()
