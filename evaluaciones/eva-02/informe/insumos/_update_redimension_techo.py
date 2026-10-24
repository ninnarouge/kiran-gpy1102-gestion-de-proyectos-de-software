# -*- coding: utf-8 -*-
"""Integra análisis completo de redimensionamiento del techo (C-05/C-06)."""
from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

DOCX = Path(__file__).resolve().parents[1] / "2.2.2-Gestion-de-cambios-Kiran v2.docx"


def set_cell_text(cell, text: str) -> None:
    if not cell.paragraphs:
        cell.text = text
        return
    p0 = cell.paragraphs[0]
    if p0.runs:
        p0.runs[0].text = text
        for run in p0.runs[1:]:
            run.text = ""
    else:
        p0.text = text
    for p in cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)


def set_para(p: Paragraph, text: str) -> None:
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.text = text


def find_para(doc: Document, contains: str = None, startswith: str = None):
    for p in doc.paragraphs:
        t = p.text or ""
        if startswith is not None and t.startswith(startswith):
            return p
        if contains is not None and contains in t:
            return p
    raise KeyError(contains or startswith)


def insert_paragraph_after(paragraph: Paragraph, text: str) -> Paragraph:
    new_p = deepcopy(paragraph._p)
    for child in list(new_p):
        if child.tag == qn("w:r"):
            new_p.remove(child)
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if paragraph.runs:
        run = new_para.add_run(text)
        src = paragraph.runs[0]
        run.bold = src.bold
        run.italic = src.italic
        if src.font.size:
            run.font.size = src.font.size
        if src.font.name:
            run.font.name = src.font.name
    else:
        new_para.add_run(text)
    return new_para


def clone_table_after(paragraph: Paragraph, template: Table) -> Table:
    new_tbl = deepcopy(template._tbl)
    paragraph._p.addnext(new_tbl)
    return Table(new_tbl, paragraph._parent)


def ensure_rows(table: Table, n: int) -> None:
    while len(table.rows) < n:
        table._tbl.append(deepcopy(table.rows[-1]._tr))
    while len(table.rows) > n:
        table.rows[-1]._tr.getparent().remove(table.rows[-1]._tr)


def fill_table(table: Table, rows: list[list[str]]) -> None:
    ensure_rows(table, len(rows))
    for i, vals in enumerate(rows):
        for cell, val in zip(table.rows[i].cells, vals):
            set_cell_text(cell, val)


def main() -> None:
    doc = Document(str(DOCX))

    # --- Intro: corregir “techo no se toca” ---
    try:
        p = find_para(doc, contains="C-05 y C-06 no suman encima")
        set_para(
            p,
            "C-05 y C-06 no suman encima de la fotografía, el segundo punto de inventario y el "
            "servicio externo en el mismo sprint: compiten por las mismas dos tarjetas por persona. "
            "Si entra uno, sale una historia equivalente. Además, C-05b implica caja nueva "
            "(Gemini + Google Cloud) que el plan inicial no tenía como partida: por eso este "
            "documento analiza la reasignación, el uso de contingencia y el redimensionamiento "
            "del techo. Un datalake completo, un especialista agéntico externo o una IA de "
            "monitoreo de flota no entran en esta ventana.",
        )
    except KeyError:
        pass

    # --- §2 evaluación: el costo SÍ puede expresar cifra nueva ---
    try:
        p = find_para(doc, contains="El costo no se expresa en una cifra nueva")
        set_para(
            p,
            "Cada cambio se examina en cronograma, alcance, costo y calidad, con el efecto "
            "positivo y el negativo. El techo de partida es US$ 20.700; si un cambio agrega "
            "partidas que no estaban en la línea base (como Gemini y Google Cloud en C-05b), "
            "el análisis debe decir si se reasigna presupuesto, si se abre contingencia con "
            "razón escrita o si se propone un techo redimensionado. No se oculta el costo bajo "
            "la frase «ya está pagado el equipo».",
        )
    except KeyError:
        pass

    # --- §4.1.1 título y párrafo ---
    h = find_para(doc, startswith="4.1.1.")
    set_para(
        h,
        "4.1.1. Desglose de gastos y redimensionamiento del presupuesto (C-05b / C-06)",
    )

    p = find_para(doc, contains="transparentes y creíbles")
    set_para(
        p,
        "C-05b y C-06 no son un ajuste gratuito sobre el plan inicial. C-06 se cubre sobre "
        "todo con horas del equipo (replanificación: 4 a 7 días-persona). C-05b agrega caja "
        "nueva: API Gemini y uso de Google Cloud para el módulo de chat, más un colchón por "
        "sobrecosto de uso, en una banda de US$ 200 a US$ 340 (propuesta de trabajo US$ 200; "
        "máximo US$ 340). Ese gasto no estaba desglosado en el acta. Además, C-05b + C-06 "
        "exigen del orden de 13 a 22 días-persona, que desplazan otro alcance del EDT si el "
        "plazo no se mueve. En zona compleja no es honesto financiar la IA vaciando la partida "
        "de horas de soporte local (~US$ 760 dentro del variable): el chatbot baja carga, no "
        "elimina al técnico de golpe. Por eso el control de cambios analiza tres escenarios y "
        "recomienda redimensionar el techo.",
    )

    # Actualizar Tabla 7b existente
    for table in doc.tables:
        blob = " ".join(c.text for r in table.rows for c in r.cells)
        if "Gemini" not in blob and "API chatbot" not in blob:
            continue
        if "Rubro" not in (table.rows[0].cells[0].text or ""):
            continue
        fill_table(
            table,
            [
                ["Rubro", "Tipo", "US$ propuesto", "US$ máximo", "Notas"],
                [
                    "API Gemini (Flash-Lite / Flash)",
                    "Variable nuevo",
                    "60",
                    "120",
                    "No estaba en el plan inicial. Precios públicos oct 2026.",
                ],
                [
                    "Google Cloud (chat, logs, cuotas)",
                    "Variable nuevo",
                    "60",
                    "100",
                    "Módulo sobre GCloud; sin GPU ni cluster.",
                ],
                [
                    "Colchón sobrecosto IA / cuotas",
                    "Reserva nombrada",
                    "80",
                    "120",
                    "Riesgo de uso mayor al estimado en piloto.",
                ],
                [
                    "HH C-05b + C-06 (replanificación)",
                    "Directo ya en sueldos",
                    "0 caja",
                    "13–22 días-persona",
                    "Desplaza alcance: si entra, sale otra historia.",
                ],
                [
                    "Subtotal caja C-05b",
                    "Incremento a financiar",
                    "200",
                    "340",
                    "Ver escenarios A/B/C de techo.",
                ],
            ],
        )
        break

    try:
        cap = find_para(doc, startswith="Tabla 7b.")
        set_para(
            cap,
            "Tabla 7b. Desglose de gastos nuevos por C-05b (Google Cloud + Gemini) y esfuerzo "
            "de C-05b/C-06. Orden de magnitud; no cotización adjudicada.",
        )
    except KeyError:
        pass

    # Insertar escenarios si no existen
    if not any("Escenario C" in (x.text or "") for x in doc.paragraphs):
        anchor = find_para(doc, startswith="Tabla 7b.")
        p = insert_paragraph_after(
            anchor,
            "Tres escenarios de financiamiento (la actividad exige elegir con transparencia):",
        )
        p = insert_paragraph_after(
            p,
            "Escenario A — Absorber sin subir el techo: se restan US$ 200 a US$ 340 del "
            "variable de horas de soporte u holgura y se crean las partidas Gemini y Google "
            "Cloud. El techo sigue en US$ 20.700. Riesgo: en zona compleja puede faltar soporte "
            "humano. Solo viable si la medición de contención del FAQ es alta desde el inicio.",
        )
        p = insert_paragraph_after(
            p,
            "Escenario B — Contingencia con razón escrita: se abre parte de los US$ 1.890 de "
            "contingencia (12 %) para el costo variable del chatbot. El techo nominal sigue en "
            "US$ 20.700, pero baja la reserva pensada para alargue de terreno u hosting. Exige "
            "acta de razón y fila en el registro de riesgos.",
        )
        p = insert_paragraph_after(
            p,
            "Escenario C — Redimensionar el techo (recomendado): se solicita ampliar el techo "
            "autorizado de US$ 20.700 a US$ 21.200 (+ US$ 500) para cubrir Gemini, Google Cloud, "
            "colchón de uso e integración, sin vaciar el presupuesto de soporte local. C-05b "
            "queda aprobado solo condicionado a esa ampliación (o, en su defecto, al Escenario B "
            "con acta). Mientras dirección no decida, no se construye el chatbot «en silencio» "
            "contra el techo viejo.",
        )

        # Tabla escenarios (clonar plantilla 5 cols = tabla 7b o tabla 3)
        template = None
        for table in doc.tables:
            if "Gemini" in " ".join(c.text for r in table.rows for c in r.cells):
                template = table
                break
        if template is None:
            template = doc.tables[2]

        new_t = clone_table_after(p, template)
        # Mover caption después: insertar párrafo y luego... la tabla quedó after p; añadir caption after table
        fill_table(
            new_t,
            [
                ["Escenario", "Techo resultante", "De dónde sale la plata", "Cuándo usarlo", "Decisión del equipo"],
                [
                    "A. Reasignar variable",
                    "US$ 20.700",
                    "Se reduce soporte/holgura variable; se crean Gemini + GCloud",
                    "Solo si se acepta menos margen de soporte humano",
                    "No recomendado como primera opción en zona compleja",
                ],
                [
                    "B. Contingencia",
                    "US$ 20.700 (reserva menor)",
                    "Parte de los US$ 1.890 con razón escrita",
                    "Si dirección no amplía techo pero autoriza el riesgo",
                    "Aceptable con acta; consume reserva de terreno",
                ],
                [
                    "C. Redimensionar techo",
                    "US$ 21.200 (+500)",
                    "Nueva autorización de dirección; partidas Gemini + GCloud en Excel",
                    "Cuando el alcance nuevo no estaba en el acta",
                    "Recomendado. C-05b condicionado a esta aprobación",
                ],
            ],
        )
        # Caption after new table
        cap_el = deepcopy(p._p)
        for child in list(cap_el):
            if child.tag == qn("w:r"):
                cap_el.remove(child)
        new_t._tbl.addnext(cap_el)
        cap_p = Paragraph(cap_el, p._parent)
        if p.runs:
            run = cap_p.add_run(
                "Tabla 7c. Escenarios de financiamiento de C-05b. El equipo recomienda el "
                "Escenario C (techo US$ 21.200)."
            )
            src = p.runs[0]
            if src.font.size:
                run.font.size = src.font.size
            if src.font.name:
                run.font.name = src.font.name
        else:
            cap_p.add_run(
                "Tabla 7c. Escenarios de financiamiento de C-05b. El equipo recomienda el "
                "Escenario C (techo US$ 21.200)."
            )

        insert_paragraph_after(
            cap_p,
            "Medición (válida en A, B o C): línea base de soporte tras el 9 de noviembre y antes "
            "del chatbot; comparar ahorro estimado de soporte con el gasto real Gemini + Google "
            "Cloud. Si el gasto de IA supera el ahorro de forma sostenida (dos revisiones), se "
            "pausa el chatbot o se abre un nuevo control de cambios, aunque el techo ya se haya "
            "ampliado.",
        )

    # --- Tabla 6 C-05 fase 1 decisión ---
    for table in doc.tables:
        if not (table.rows[0].cells[0].text or "").startswith("Id"):
            continue
        if "Decisión" not in (table.rows[0].cells[1].text or ""):
            continue
        for row in table.rows:
            if (row.cells[0].text or "").startswith("C-05 fase 1"):
                set_cell_text(
                    row.cells[1],
                    "Aprobada condicionada a C-06, WIP, FAQ curado y financiamiento explícito "
                    "(Escenario C: techo US$ 21.200, o B con acta de contingencia). Fases 2–3 no. "
                    "El bot no cierra ni cambia estados solo.",
                )
                set_cell_text(
                    row.cells[4],
                    "Sí, a la dirección: ampliar techo a US$ 21.200, o abrir contingencia, o "
                    "rechazar C-05b. No construir contra el techo viejo sin esa decisión.",
                )

    # --- Tabla 10 C-05b ---
    for table in doc.tables:
        if not (table.rows[0].cells[0].text or "").startswith("Id"):
            continue
        if len(table.rows[0].cells) < 6:
            continue
        for row in table.rows:
            if (row.cells[0].text or "").startswith("C-05b"):
                set_cell_text(
                    row.cells[5],
                    "Abierto condicionado a C-06, WIP, FAQ curado y a la decisión de "
                    "financiamiento: recomendado Escenario C (techo US$ 21.200, +500 para "
                    "Gemini + Google Cloud + colchón). Alternativa B: contingencia con acta. "
                    "Sin esa decisión no se construye. Medición ahorro soporte vs gasto IA "
                    "obligatoria. Escalamiento a dirección si proponen absorber solo vaciando "
                    "soporte humano (Escenario A) o adelantar monitoreo con IA.",
                )

    # --- Tabla 7 Excel C-05 ---
    for table in doc.tables:
        if not (table.rows[0].cells[0].text or "").startswith("Decisión"):
            continue
        for row in table.rows:
            if "C-05" in (row.cells[0].text or "") and "fase 1" in (row.cells[0].text or ""):
                set_cell_text(
                    row.cells[4],
                    "Ari crea partidas Gemini y Google Cloud solo si hay Escenario C (techo "
                    "21.200) o B (contingencia con acta). Registra gasto real y lo compara con "
                    "ahorro de soporte (línea base). Sin decisión de financiamiento: sin "
                    "despliegue. Detalle en Tablas 7b y 7c.",
                )

    # --- Tabla 11 ---
    t11 = doc.tables[-1]
    if "Pregunta" in (t11.rows[0].cells[0].text or ""):
        set_cell_text(
            t11.rows[1].cells[1],
            "La prioridad y la nota en el ticket (C-01b) si reemplazan historia. La regla "
            "C-04. C-06 condicionada. C-05b (chatbot FAQ + canal en Google Cloud/Gemini) "
            "condicionada además a financiamiento explícito: techo redimensionado a US$ "
            "21.200 (recomendado) o contingencia con acta.",
        )
        # ¿Qué no se toca? — actualizar
        for row in t11.rows:
            if "no se toca" in (row.cells[0].text or "").lower() or "Qué no se toca" in (
                row.cells[0].text or ""
            ):
                set_cell_text(
                    row.cells[1],
                    "Las fechas de los cortes y los cuatro roles. El techo de US$ 20.700 es el "
                    "de partida; si se aprueba C-05b en Escenario C, el techo autorizado pasa a "
                    "US$ 21.200 por control de cambios (no por ajuste informal). La contingencia "
                    "sigue sin usarse salvo razón escrita (Escenario B).",
                )

    # --- E4 ---
    try:
        e4 = find_para(doc, startswith="E4.")
        set_para(
            e4,
            "E4. Techo y gasto comprometido: desviación acumulada frente al techo vigente "
            "(US$ 20.700 de partida; US$ 21.200 si dirección aprueba el Escenario C). Freno si "
            "supera el 10 %. Dueño: Ari. Si se rompe: detener gasto no comprometido (Gemini/"
            "Google Cloud). C-05b no se gasta contra el techo viejo sin escenario A, B o C "
            "aprobado.",
        )
    except KeyError:
        pass

    # --- Cierre ---
    try:
        cierre = find_para(doc, contains="Proyecto: Kiran. Organización")
        t = cierre.text or ""
        if "21.200" not in t:
            set_para(
                cierre,
                t.rstrip()
                + " Presupuesto: techo de partida US$ 20.700; propuesta de redimensionamiento a "
                "US$ 21.200 si se aprueba C-05b (Escenario C), con desglose Gemini + Google Cloud "
                "en Tablas 7b y 7c.",
            )
    except KeyError:
        pass

    # --- §3.1 C-05 párrafo ---
    try:
        p = find_para(doc, startswith="C-05 se aprueba solo por fases")
        t = p.text or ""
        if "21.200" not in t:
            set_para(
                p,
                t.rstrip()
                + " El financiamiento de la fase 1 exige decisión explícita: Escenario C "
                "(techo US$ 21.200, recomendado), Escenario B (contingencia con acta) o "
                "rechazo. No se construye asumiendo que «entra en los 20.700» sin desglose.",
            )
    except KeyError:
        pass

    doc.save(str(DOCX))
    print("OK", DOCX)


if __name__ == "__main__":
    main()
