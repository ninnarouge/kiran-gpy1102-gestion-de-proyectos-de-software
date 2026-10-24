# -*- coding: utf-8 -*-
"""Fija stack GCloud+Gemini y medición gasto API vs ahorro de soporte."""
from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

DOCX = Path(__file__).resolve().parents[1] / "2.2.2-Gestion-de-cambios-Kiran v2.docx"


def set_cell_text(cell, text: str) -> None:
    if not cell.paragraphs:
        cell.text = text
        return
    p0 = cell.paragraphs[0]
    if p0.runs:
        p0.runs[0].text = text
        for run in p0.runs[1:]:
            run.text = ""
    else:
        p0.text = text
    for p in cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)


def set_para(p: Paragraph, text: str) -> None:
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.text = text


def find_para(doc: Document, contains: str = None, startswith: str = None):
    for p in doc.paragraphs:
        t = p.text or ""
        if startswith is not None and t.startswith(startswith):
            return p
        if contains is not None and contains in t:
            return p
    raise KeyError(contains or startswith)


def insert_paragraph_after(paragraph: Paragraph, text: str) -> Paragraph:
    new_p = deepcopy(paragraph._p)
    for child in list(new_p):
        if child.tag == qn("w:r"):
            new_p.remove(child)
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if paragraph.runs:
        run = new_para.add_run(text)
        src = paragraph.runs[0]
        run.bold = src.bold
        run.italic = src.italic
        if src.font.size:
            run.font.size = src.font.size
        if src.font.name:
            run.font.name = src.font.name
    else:
        new_para.add_run(text)
    return new_para


def main() -> None:
    doc = Document(str(DOCX))

    # §4.1.1 párrafo principal
    p = find_para(doc, contains="transparentes y creíbles")
    set_para(
        p,
        "Para que C-05b y C-06 sean transparentes y creíbles, el equipo fija un stack de "
        "referencia: Google Cloud (despliegue del módulo de chat, por ejemplo Cloud Run o "
        "Firebase App Hosting sobre el ambiente del piloto) y Gemini en gama Flash / "
        "Flash-Lite vía API gestionada (precios públicos revisados en octubre de 2026; p. ej. "
        "Flash-Lite en el orden de US$ 0,10 por 1M de tokens de entrada y US$ 0,40 por 1M de "
        "salida, u opciones Flash según necesidad). El costo incremental de caja de la fase 1 "
        "se estima en una banda de US$ 120 a US$ 220 dentro del costo variable del piloto "
        "(consumo Gemini + margen de uso en Google Cloud), con tope duro de API Gemini de "
        "US$ 120 en Excel. Ese monto no sube el techo de US$ 20.700 ni abre la contingencia "
        "del 12 %. El esfuerzo (13 a 22 días-persona con C-06) se cubre replanificando el "
        "equipo de cuatro. No se contrata especialista agéntico externo en esta ventana. "
        "La justificación económica no se declara a ojo: se mide. Antes del chatbot se toma "
        "una línea base de consultas de soporte (horas o cantidad fuera del tablero); después "
        "se comparan el gasto real de Gemini/Google Cloud (factura o consola) con el soporte "
        "evitado. Si el gasto de IA supera el ahorro medido de soporte de forma sostenida, se "
        "pausa el chatbot o se abre un control de cambios. Las cifras son orden de magnitud "
        "con tráfico de piloto; no son cotización adjudicada.",
    )

    # Actualizar Tabla 7b (buscar la que tiene API chatbot)
    for table in doc.tables:
        blob = " ".join(c.text for r in table.rows for c in r.cells)
        if "API chatbot" not in blob and "Gemini" not in blob:
            continue
        if "US$ propuesto" not in blob and "Rubro" not in table.rows[0].cells[0].text:
            continue
        # Es la 7b
        rows_data = [
            [
                "Rubro",
                "Tipo de costo",
                "US$ propuesto",
                "US$ máximo",
                "Supuesto / control",
            ],
            [
                "API Gemini (Flash-Lite o Flash en Google AI / Vertex)",
                "Variable y directo",
                "60",
                "120",
                "Precios públicos Gemini (oct 2026). Tráfico piloto. Ari tope en Excel; alerta al 80 %.",
            ],
            [
                "Google Cloud (Cloud Run / App Hosting, logs, cuota)",
                "Variable y directo",
                "60",
                "100",
                "Módulo chat en GCloud sobre el ambiente del piloto; sin GPU ni cluster.",
            ],
            [
                "HH equipo (C-05b + C-06)",
                "Directo ya en el techo",
                "0 (caja)",
                "13 a 22 días-persona",
                "Replanificación de historias. Sin contratar especialista agéntico.",
            ],
            [
                "Ahorro de soporte (a medir, no presupuesto de gasto)",
                "Beneficio operativo",
                "Baseline post 9 nov",
                "Meta: ≥ gasto Gemini+GCloud del período",
                "Horas o consultas evitadas × costo hora soporte local. Dueña medición: Skarlett + Ari.",
            ],
            [
                "Total caja nueva de C-05b",
                "Variable dentro del techo",
                "120",
                "220",
                "Si gasto IA > ahorro soporte medido de forma sostenida → pausar o nuevo CC. Contingencia cerrada.",
            ],
        ]
        # Asegurar 6 filas
        while len(table.rows) < 6:
            tr = deepcopy(table.rows[-1]._tr)
            table._tbl.append(tr)
        while len(table.rows) > 6:
            table.rows[-1]._tr.getparent().remove(table.rows[-1]._tr)
        for i, vals in enumerate(rows_data):
            for cell, val in zip(table.rows[i].cells, vals):
                set_cell_text(cell, val)
        break

    # Caption 7b
    try:
        cap = find_para(doc, startswith="Tabla 7b.")
        set_para(
            cap,
            "Tabla 7b. Estimación y medición de costo C-05b/C-06 con Google Cloud y Gemini "
            "(orden de magnitud; no cotización adjudicada).",
        )
    except KeyError:
        pass

    # Tabla 7 fila C-05 excel
    for table in doc.tables:
        if not (table.rows[0].cells[0].text or "").startswith("Decisión"):
            continue
        for row in table.rows:
            if "C-05" in (row.cells[0].text or "") and "fase 1" in (row.cells[0].text or ""):
                set_cell_text(
                    row.cells[4],
                    "Ari registra en Excel: gasto Gemini (tope US$ 120), uso Google Cloud "
                    "(hasta US$ 100) y, en paralelo, horas o consultas de soporte evitadas "
                    "frente a la línea base. Banda total US$ 120–220. Si el gasto de IA supera "
                    "el ahorro de soporte medido de forma sostenida, se pausa el chatbot. "
                    "Contingencia del 12 % cerrada. Detalle en Tabla 7b.",
                )

    # Medición en §4.4 — insertar si no existe
    if not any("línea base de soporte" in (p.text or "") for p in doc.paragraphs):
        anchor = find_para(doc, contains="Los operativos O1 a O4 sostienen")
        insert_paragraph_after(
            anchor,
            "Medición económica de C-05b (compensa soporte vs Gemini/Google Cloud). "
            "1) Línea base: durante al menos una semana con tablero usable y antes de activar "
            "el chatbot, Skarlett y el soporte anotan consultas fuera del sistema (cantidad y "
            "minutos). 2) En operación: Ari lee de la consola de Google Cloud el gasto de "
            "Gemini y del servicio de chat; Skarlett anota contención FAQ (O6) y consultas "
            "escaladas. 3) Regla: el período se considera sostenible si el costo estimado de "
            "soporte evitado es mayor o igual al gasto Gemini + Google Cloud del mismo "
            "período; si el gasto de IA supera ese ahorro de forma sostenida (dos revisiones "
            "seguidas), se pausa el chatbot o se abre un control de cambios. No se declara "
            "ahorro en dólares en el acta hasta tener esa línea base.",
        )

    # O6 submétricas — ampliar si existe párrafo O6
    try:
        o6 = find_para(doc, startswith="O6.")
        t = o6.text or ""
        if "Gemini" not in t:
            set_para(
                o6,
                t.rstrip()
                + " Incluye submétricas económicas: gasto Gemini/Google Cloud del período "
                "(Ari) frente a consultas o minutos de soporte evitados respecto de la línea base.",
            )
    except KeyError:
        pass

    doc.save(str(DOCX))
    print("OK", DOCX)


if __name__ == "__main__":
    main()
