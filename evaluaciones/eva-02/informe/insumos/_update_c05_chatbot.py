# -*- coding: utf-8 -*-
"""Reencuadra C-05: chatbot FAQs + canal de soporte (ahorro de personal)."""
from __future__ import annotations

from pathlib import Path

from docx import Document

DOCX = Path(__file__).resolve().parents[1] / "2.2.2-Gestion-de-cambios-Kiran.docx"

# Texto canónico C-05 (Tabla 2)
C05_SITUACION = (
    "Ejemplo de solicitud. La dirección pide un chatbot en la plataforma para los usuarios "
    "del piloto: preguntas frecuentes (FAQ) y un canal para comunicarse por temas de soporte "
    "y asistencia, con el fin de ahorrar contrato de personal de soporte. En este plan la "
    "formulación amplia (app ciudadana masiva o reemplazo total del técnico) no se aprueba: "
    "solo se contempla la fase 1 (FAQ publicadas en el chat más escalamiento a ticket o a "
    "persona de soporte cuando el FAQ no alcanza), tras un tablero usable, sin datos de "
    "hogares en un modelo externo y sin sustituir los estados del kit."
)

REPLACEMENTS = [
    # Intro anti-desborde
    (
        "Un datalake completo, un modelo hospedado por el equipo o una IA para hogares y donantes no "
        "entran en esta ventana: o se cierran, o quedan como solicitud futura fuera del piloto.",
        "Un datalake completo, un modelo hospedado por el equipo, una app ciudadana masiva o un "
        "chatbot que reemplace por completo al soporte contratado no entran en esta ventana: o se "
        "cierran, o quedan como solicitud futura fuera del piloto. C-05 fase 1 es autoservicio de "
        "FAQ y canal de asistencia para usuarios de la plataforma, no un segundo producto.",
    ),
    # §1.1 párrafo C-05
    (
        "C-05 se rotula como ejemplo de solicitud porque concreta el tipo de cambio de productividad "
        "con IA que el piloto puede recibir. No es un copiloto ya construido. Lo que amerita registro "
        "es la ampliación de superficie de producto y el costo variable de una API de modelo. La "
        "medida factible, si cabe, es solo la fase 1 de soporte.",
        "C-05 se rotula como ejemplo de solicitud porque concreta el tipo de cambio de soporte al "
        "usuario con IA que el piloto puede recibir: un chatbot con FAQ y canal de asistencia para "
        "bajar el gasto en personal de soporte contratado. No es una app ciudadana ni el reemplazo "
        "total del técnico. Lo que amerita registro es la nueva superficie en la plataforma y el "
        "costo variable de una API de modelo. La medida factible, si cabe, es solo la fase 1 "
        "(FAQ más escalamiento humano).",
    ),
    # §2.5 título
    (
        "2.5. C-05. Asistente de IA de uso interno",
        "2.5. C-05. Chatbot de FAQ y canal de soporte",
    ),
    # §2.5 cuerpos
    (
        "En los términos pedidos (copiloto para tickets, reportes y decisiones), el cambio toca "
        "integración, costo variable de API, privacidad y calidad de las sugerencias. Nicolás "
        "sostiene la integración. Skarlett valida la calidad de lo sugerido. Ari registra el "
        "tope de gasto. Giannina sostiene el control de cambios. Un modelo propio, un fine-tuning "
        "o un datalake para entrenar no están en el plan.",
        "En los términos pedidos (chatbot para usuarios con FAQ y comunicación de soporte, para "
        "ahorrar contrato de personal), el cambio toca la interfaz de la plataforma, integración "
        "con una API de modelo o motor de FAQ, privacidad, calidad de las respuestas y el flujo "
        "hacia el ticket. Nicolás sostiene la integración. Giannina la interfaz del chat. "
        "Skarlett valida la calidad de las respuestas y el FAQ. Ari registra el tope de gasto. "
        "Un modelo propio, un fine-tuning, un datalake para entrenar o una app para cada hogar "
        "no están en el plan.",
    ),
    (
        "El efecto positivo acotado es aliviar al técnico a 0,4 FTE: menos tiempo para entender un "
        "ticket y elegir el siguiente, si la sugerencia se usa y se corrige. El efecto negativo "
        "de la formulación completa es ampliar el alcance, abrir gasto no comprometido y "
        "priorizar mal por una alucinación. La fase 1 no cambia el estado del kit ni redacta el "
        "reporte a donantes. Indicadores amenazados: E1, E4 y E5. Operativos a monitorear: O1, "
        "O4 y O6, con submetricas de adopción y rechazo de sugerencias.",
        "El efecto positivo acotado es reducir consultas repetidas al personal de soporte "
        "contratado: el usuario resuelve con FAQ o deja un mensaje que escala a ticket. El "
        "efecto negativo de la formulación completa (reemplazo total del soporte o app ciudadana) "
        "es ampliar el alcance, abrir gasto no comprometido y responder mal por una alucinación. "
        "La fase 1 no cambia el estado del kit, no cierra tickets sola y no redacta el reporte a "
        "donantes. Indicadores amenazados: E1, E4 y E5. Operativos a monitorear: O1, O4 y O6, "
        "con submétricas de contención del chatbot (% resuelto por FAQ sin escalar) y de "
        "escalamiento a persona.",
    ),
    (
        "presión de esfuerzo hacia el 9 de noviembre y más líneas cambiadas antes del 4 de "
        "diciembre, con riesgo sobre O1 y, si persiste, sobre E1. El orden de adopción es "
        "C-06 antes que la IA fase 1: la IA no adelanta a un tablero que el soporte aún no usa.",
        "presión de esfuerzo hacia el 9 de noviembre y más líneas cambiadas antes del 4 de "
        "diciembre, con riesgo sobre O1 y, si persiste, sobre E1. El orden de adopción es "
        "C-06 (cola visible para quien atiende) antes que el chatbot fase 1: el canal de "
        "asistencia no adelanta a un tablero y una cola que el soporte aún no usa.",
    ),
    (
        "del hosting variable, sin competir con C-06. C-05 fase 1 y C-06 solo entran si reemplazan "
        "historia y respetan dos tarjetas por persona. Hay que evitar acumular superficie de "
        "producto sin capacidad de integración y prueba.",
        "del hosting variable, sin competir con C-06. C-05 fase 1 (chatbot FAQ + canal) y C-06 "
        "solo entran si reemplazan historia y respetan dos tarjetas por persona. Hay que evitar "
        "acumular superficie de producto sin capacidad de integración y prueba.",
    ),
    # Tabla 3 / riesgos / soluciones — textos largos
    (
        "C-05 se aprueba solo por fases. La fase 1 (soporte) queda condicionada a que C-06 y, si "
        "aplica, C-01b estén en marcha o en el mismo plan de sprint, a que quepa en dos tarjetas "
        "por persona y a un tope de llamadas API en Excel. La fase 2 (borradores internos de cierre) "
        "solo post 4 de diciembre o con un control de cambios nuevo. La fase 3 (copiloto amplio, "
        "priorización automática, IA a hogares o donantes, modelo propio) se rechaza entre "
        "septiembre y diciembre de 2026. No se aprueba un copiloto único «todo en uno».",
        "C-05 se aprueba solo por fases. La fase 1 (chatbot con FAQ y canal de asistencia que "
        "escala a ticket o a persona) queda condicionada a que C-06 esté en marcha o en el mismo "
        "plan de sprint, a que quepa en dos tarjetas por persona y a un tope de llamadas API en "
        "Excel. El FAQ lo cura el equipo; el chatbot no cierra ni cambia estados solo. La fase 2 "
        "(ampliar el corpus FAQ y mejorar el enrutamiento del canal) solo post 4 de diciembre o "
        "con un control de cambios nuevo. La fase 3 (reemplazo total del personal de soporte, "
        "app ciudadana masiva, modelo propio) se rechaza entre septiembre y diciembre de 2026. "
        "No se aprueba un chatbot que sustituya de golpe al soporte contratado.",
    ),
    (
        "C-06 se aprueba en forma condicionada: vistas de cola o filtro en el tablero hacia el 9 de "
        "noviembre, sin aprendizaje automático, midiendo menos consultas informales. Es prerrequisito "
        "de la IA fase 1. Si no cabe en el límite de trabajo en curso, no entra.",
        "C-06 se aprueba en forma condicionada: vistas de cola o filtro en el tablero hacia el 9 de "
        "noviembre, sin aprendizaje automático, midiendo menos consultas informales. Es prerrequisito "
        "del chatbot fase 1: lo que escale desde el chat debe poder atenderse en la cola. Si no cabe "
        "en el límite de trabajo en curso, no entra.",
    ),
    # §4.4 O6 submétricas
    (
        "después del 9 de noviembre. Submetricas de C-05: porcentaje de tickets con resumen o "
        "checklist usado y porcentaje de sugerencias rechazadas; rechazo alto congela la fase 2. "
        "Dueños: Giannina y el soporte; en IA, Nicolás y Skarlett.",
        "después del 9 de noviembre. Submétricas de C-05: porcentaje de consultas del chatbot "
        "contenidas por FAQ sin escalar; porcentaje que escala a ticket o a persona; calidad "
        "percibida del FAQ (revisión de Skarlett). Contención baja o alucinaciones altas congelan "
        "la fase 2. Dueños: Giannina y el soporte; en el chatbot, Nicolás y Skarlett.",
    ),
    (
        "Los operativos O1 a O4 sostienen ritmo y calidad del equipo. O5 a O7 miden si el piloto "
        "opera con datos creíbles y soporte atendible. E1 a E3 consolidan eso en cortes y EDT; E4 "
        "y E5 en techo y control de cambios.",
        "Los operativos O1 a O4 sostienen ritmo y calidad del equipo. O5 a O7 miden si el piloto "
        "opera con datos creíbles y soporte atendible, incluido el efecto del chatbot sobre la "
        "carga de personal. E1 a E3 consolidan eso en cortes y EDT; E4 y E5 en techo y control de cambios.",
    ),
    (
        " Registro ampliado con C-05a (cerrado), C-05b (abierto condicionado) y C-06 (abierto "
        "condicionado).",
        " Registro ampliado con C-05a (chatbot amplio / app ciudadana / reemplazo total: cerrado), "
        "C-05b (FAQ + canal fase 1: abierto condicionado) y C-06 (abierto condicionado).",
    ),
]


def set_cell_text(cell, text: str) -> None:
    if not cell.paragraphs:
        cell.text = text
        return
    p0 = cell.paragraphs[0]
    if p0.runs:
        p0.runs[0].text = text
        for run in p0.runs[1:]:
            run.text = ""
    else:
        p0.text = text
    for p in cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)


def replace_text_in_paragraphs(doc: Document, old: str, new: str) -> int:
    n = 0
    for p in doc.paragraphs:
        t = p.text or ""
        if old in t:
            full = t.replace(old, new)
            if p.runs:
                p.runs[0].text = full
                for r in p.runs[1:]:
                    r.text = ""
            else:
                p.text = full
            n += 1
    return n


def replace_in_tables(doc: Document, old: str, new: str) -> int:
    n = 0
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                t = cell.text or ""
                if old in t:
                    set_cell_text(cell, t.replace(old, new))
                    n += 1
    return n


def find_row(table, first_cell_startswith: str):
    for row in table.rows:
        if (row.cells[0].text or "").startswith(first_cell_startswith):
            return row
    return None


def main() -> None:
    doc = Document(str(DOCX))
    tables = doc.tables
    t2, t3, t4, t5, t6, t7, t8, t9, t10, t11 = (
        tables[1],
        tables[2],
        tables[3],
        tables[4],
        tables[5],
        tables[6],
        tables[7],
        tables[8],
        tables[9],
        tables[10],
    )

    # Tabla 2 — fila C-05
    row = find_row(t2, "C-05")
    assert row is not None
    set_cell_text(row.cells[1], C05_SITUACION)
    set_cell_text(row.cells[2], "Dirección / equipo (ahorro de personal de soporte)")
    set_cell_text(
        row.cells[3],
        "9 de noviembre (depende del tablero usable) y 4 de diciembre (tickets y canal). "
        "No la ruta crítica de inventario.",
    )

    # Tabla 3
    row = find_row(t3, "C-05 completo")
    if row:
        set_cell_text(row.cells[0], "C-05 completo (app ciudadana / reemplazo total / datalake)")
        set_cell_text(
            row.cells[1],
            "Alto. Abre app masiva o elimina el soporte humano del piloto sin capacidad de integración.",
        )
        set_cell_text(
            row.cells[2],
            "Gasto variable alto y posible contingencia. Rompe E4 si supera el 10 % de desviación.",
        )
        set_cell_text(row.cells[3], "Rompe o desplaza el 4 de diciembre.")
        set_cell_text(row.cells[4], "No")

    row = find_row(t3, "C-05 fase 1")
    if row:
        set_cell_text(row.cells[0], "C-05 fase 1 (chatbot FAQ + canal)")
        set_cell_text(
            row.cells[1],
            "Medio. Chat en la plataforma con FAQ y escalamiento a ticket o persona; no cierra solo.",
        )
        set_cell_text(
            row.cells[2],
            "Horas replanificadas más sobre API o hosting del chat dentro del costo variable. "
            "Contingencia cerrada salvo razón escrita. Objetivo: menos horas de soporte contratado.",
        )
        set_cell_text(
            row.cells[3],
            "Post 9 de noviembre; no en la semana del 5 al 9 de diciembre.",
        )
        set_cell_text(
            row.cells[4],
            "Condicionado a C-06, cupo de sprint, tope de API y FAQ curado por el equipo.",
        )

    # Tabla 4 — riesgos IA
    for row in t4.rows:
        r0 = row.cells[0].text or ""
        if r0.startswith("La IA amplía"):
            set_cell_text(
                row.cells[0],
                "El chatbot amplía el alcance sin control.",
            )
            set_cell_text(
                row.cells[1],
                "Se construye app ciudadana, se elimina el soporte contratado de golpe o el bot "
                "cierra tickets y cambia estados solo.",
            )
            set_cell_text(
                row.cells[2],
                "Solo fase 1: FAQ más escalamiento humano. Un salto de fase es solicitud nueva. "
                "El análisis de impacto lo revisa un rol senior (Ari o Nicolás), no un junior en "
                "planificación ni en control de cambios.",
            )
        if "datalake come" in r0 or r0.startswith("El costo de API"):
            set_cell_text(
                row.cells[1],
                "Se contrata un lago completo, un modelo propio o se abre la contingencia sin acta "
                "por el chatbot.",
            )
            set_cell_text(
                row.cells[2],
                "El datalake completo y la app ciudadana se rechazan o se difieren. Ari pone tope "
                "de API del chatbot en Excel. Si la desviación supera el 10 %, se pausa el chatbot "
                "antes que el corte.",
            )

    # Tabla 5
    row = find_row(t5, "C-05 fase 1")
    if row:
        set_cell_text(row.cells[0], "C-05 fase 1, chatbot FAQ + canal")
        set_cell_text(
            row.cells[1],
            "Sí para bajar carga de personal de soporte, no antes que la cola de atención.",
        )
        set_cell_text(row.cells[2], "No antes del tablero usable.")
        set_cell_text(row.cells[3], "Planificar después del 9 de noviembre.")
        set_cell_text(
            row.cells[4],
            "Después de C-06. Condicionada al cupo, al tope de API y al FAQ curado.",
        )

    row = find_row(t5, "C-05 fases")
    if row:
        set_cell_text(
            row.cells[0],
            "C-05 fases 2–3 / app ciudadana / datalake completo",
        )
        set_cell_text(row.cells[1], "No en esta ventana.")
        set_cell_text(row.cells[2], "No.")
        set_cell_text(
            row.cells[3],
            "Sacar del camino o planificar sin fecha de construcción.",
        )
        set_cell_text(
            row.cells[4],
            "Rechazado o diferido. Sin pasos en sep–dic 2026.",
        )

    # Tabla 6
    row = find_row(t6, "C-05 fase 1")
    if row:
        set_cell_text(
            row.cells[1],
            "Aprobada condicionada a C-06, al límite de dos tarjetas, al tope de API y a un FAQ "
            "curado. Las fases 2 y 3 no. El bot no cierra ni cambia estados solo.",
        )
        set_cell_text(
            row.cells[2],
            "Ofrece autoservicio y canal de asistencia a usuarios de la plataforma para ahorrar "
            "personal de soporte, sin abrir app ciudadana ni reporte automático a donantes.",
        )
        set_cell_text(row.cells[3], "Abierto, condicionado")
        set_cell_text(
            row.cells[4],
            "Sí, a la dirección, si proponen mover el 4 de diciembre, abrir la contingencia por "
            "API o reemplazar de golpe al soporte contratado.",
        )

    row = find_row(t6, "C-05 completo")
    if row:
        set_cell_text(row.cells[0], "C-05 completo / app ciudadana / datalake")
        set_cell_text(
            row.cells[1],
            "No se aprueba en esta ventana. Se cierra esa formulación.",
        )
        set_cell_text(
            row.cells[2],
            "El piloto no abre una app masiva ni un segundo producto de datos; el soporte humano "
            "sigue existiendo para lo que el FAQ no resuelve.",
        )
        set_cell_text(row.cells[3], "Cerrado")
        set_cell_text(
            row.cells[4],
            "Sí, a la dirección, si se insiste en app ciudadana, lago o modelo propio.",
        )

    # Tabla 7
    row = find_row(t7, "C-05 fase 1")
    if row:
        set_cell_text(row.cells[0], "C-05 fase 1, chatbot condicionada")
        set_cell_text(
            row.cells[1],
            "Se agrega chatbot en la plataforma: corpus FAQ, ventana de chat y escalamiento a "
            "ticket o a persona de soporte. Esfuerzo orientativo: cuatro a seis días-persona de "
            "integración y dos a tres de calidad del FAQ. No cambia estados del kit solo.",
        )
        set_cell_text(
            row.cells[2],
            "Las tareas quedan bajo el paquete de tickets y soporte. La predecesora es el tablero "
            "usable. No se construye en la semana del 5 al 9 de diciembre.",
        )
        set_cell_text(
            row.cells[3],
            "Misma regla de reemplazo de historia. No se suma a C-06 en la misma persona sin sacar trabajo.",
        )
        set_cell_text(
            row.cells[4],
            "Ari registra un subítem de sobre API del chatbot dentro del costo variable. Tope "
            "mensual escrito. Contingencia del 12 % cerrada. Desviación sobre el 10 %: se pausa "
            "el chatbot. El ahorro esperado es de horas de soporte contratado, no un segundo techo.",
        )

    row = find_row(t7, "C-05 amplio")
    if row:
        set_cell_text(row.cells[0], "C-05 amplio / app ciudadana / datalake, cerrado")
        set_cell_text(row.cells[1], "El alcance del piloto no cambia.")
        set_cell_text(row.cells[2], "No se inserta en el cronograma de esta ventana.")
        set_cell_text(row.cells[3], "Sin tarjeta de construcción.")
        set_cell_text(row.cells[4], "Sin gasto. Queda cerrado o como solicitud post-piloto.")

    # Tabla 8 — pasos C-05
    for row in t8.rows:
        p0 = row.cells[0].text or ""
        if p0.startswith("Tras el tablero usable"):
            set_cell_text(
                row.cells[0],
                "Tras el tablero usable, decidir si entra C-05 fase 1 (chatbot FAQ + canal): cupo "
                "de dos tarjetas, tope de API de Ari, FAQ curado y Review con quien da soporte.",
            )
            set_cell_text(
                row.cells[1],
                "Giannina Guerrero (control de cambios e interfaz del chat); Nicolás Barra "
                "(integración); Skarlett Tropan (calidad del FAQ y respuestas); Ari Araya (gasto)",
            )
        if "porcentaje de tickets con resumen" in p0 or p0.startswith("Medir, bajo el indicador O6"):
            set_cell_text(
                row.cells[0],
                "Medir, bajo el indicador O6, el porcentaje de consultas contenidas por FAQ sin "
                "escalar y el porcentaje que escala a ticket o a persona. Si la contención es baja "
                "o hay alucinaciones, congelar la fase 2.",
            )
        if p0.startswith("No abrir la fase 2"):
            set_cell_text(
                row.cells[0],
                "No abrir la fase 2 del chatbot, ni app ciudadana, ni un datalake completo sin un "
                "control de cambios nuevo y revisión de un rol senior.",
            )

    # Tabla 9
    set_cell_text(
        t9.rows[2].cells[3],
        "Puede entrar C-06 (cola o filtros), solo a cambio de otra historia. No entra un datalake "
        "ni una app ciudadana. No entra C-05 fase 1 (chatbot) antes del tablero usable.",
    )
    set_cell_text(
        t9.rows[3].cells[3],
        "Puede entrar la prioridad y la nota, y C-05 fase 1 (chatbot FAQ + canal de asistencia) si "
        "cumplió condiciones, solo a cambio de otra historia. No entra el archivo. No entra un "
        "servicio externo. No entra la fase 2 del chatbot ni un datalake ni una app ciudadana.",
    )

    # Tabla 10
    row = find_row(t10, "C-05a")
    if row:
        set_cell_text(
            row.cells[1],
            "Requisito. Chatbot amplio, app ciudadana o reemplazo total del soporte contratado; o datalake completo.",
        )
        set_cell_text(
            row.cells[5],
            "No aprobado entre septiembre y diciembre de 2026. Queda fuera de la EDT y del techo. "
            "Escalamiento a la dirección si se insiste.",
        )
    row = find_row(t10, "C-05b")
    if row:
        set_cell_text(
            row.cells[1],
            "Requisito. Chatbot fase 1: FAQ en la plataforma y canal de asistencia con escalamiento a ticket o persona.",
        )
        set_cell_text(
            row.cells[5],
            "Aprobado solo si C-06 está en marcha, cabe en dos tarjetas, existe tope de API y el "
            "FAQ está curado. Sin datos de hogares en un modelo externo. El bot no cierra ni "
            "cambia estados solo. Objetivo: ahorrar horas de soporte contratado. Escalamiento si "
            "proponen mover el 4 de diciembre o reemplazar de golpe al personal de soporte.",
        )

    # Tabla 11
    set_cell_text(
        t11.rows[1].cells[1],
        "La prioridad y la nota de texto en el ticket, y solo si reemplazan una historia "
        "equivalente sin mover el 4 de diciembre. La regla de reasignación cuando un bloqueo "
        "supera cinco días. La priorización visual del tablero (C-06) y el chatbot fase 1 de "
        "FAQ y canal de asistencia (C-05b), ambas condicionadas al reemplazo de historia, al "
        "límite de dos tarjetas y, en el chatbot, al tope de API, al FAQ curado y al prerrequisito de cola.",
    )
    set_cell_text(
        t11.rows[2].cells[1],
        "La fotografía del ticket, el segundo punto de inventario dentro de esta ventana, el "
        "servicio externo del reporte, el chatbot que reemplaza de golpe al soporte contratado, "
        "la app ciudadana masiva, el datalake o lakehouse completo, el modelo propio y las fases "
        "2 y 3 del chatbot en esta ventana.",
    )
    set_cell_text(
        t11.rows[3].cells[1],
        "Kiran no instala paneles. Ninguna solicitud de este plan modifica ese límite. Tampoco "
        "abre una app ciudadana masiva ni una plataforma analítica aparte del piloto. El chatbot "
        "fase 1 atiende a usuarios de la plataforma del piloto; no convierte a cada hogar en "
        "usuario de una app nueva.",
    )

    total = 0
    for old, new in REPLACEMENTS:
        total += replace_text_in_paragraphs(doc, old, new)
        total += replace_in_tables(doc, old, new)

    # Reemplazos cortos residuales en tablas/párrafos
    short = [
        ("asistente de IA de uso interno", "chatbot de FAQ y canal de soporte"),
        ("IA fase 1 de soporte", "chatbot fase 1 (FAQ + canal)"),
        ("la IA fase 1", "el chatbot fase 1"),
        ("de la IA fase 1", "del chatbot fase 1"),
        ("copiloto de IA amplio", "chatbot amplio o app ciudadana"),
        ("sobre API de IA", "sobre API del chatbot"),
        ("API de IA", "API del chatbot"),
        ("se pausa la IA", "se pausa el chatbot"),
        ("se pausa IA", "se pausa el chatbot"),
        ("calidad de sugerencias", "calidad del FAQ y respuestas"),
        ("calidad de lo sugerido", "calidad del FAQ y de las respuestas"),
    ]
    for old, new in short:
        total += replace_text_in_paragraphs(doc, old, new)
        total += replace_in_tables(doc, old, new)

    doc.save(str(DOCX))
    print("OK replacements~", total, DOCX)


if __name__ == "__main__":
    main()
