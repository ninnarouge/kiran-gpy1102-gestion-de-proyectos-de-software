# -*- coding: utf-8 -*-
"""Agrega dimensión CAPEX / OPEX al análisis presupuestario C-05/C-06."""
from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

DOCX = Path(__file__).resolve().parents[1] / "2.2.2-Gestion-de-cambios-Kiran v2.docx"


def set_cell_text(cell, text: str) -> None:
    if not cell.paragraphs:
        cell.text = text
        return
    p0 = cell.paragraphs[0]
    if p0.runs:
        p0.runs[0].text = text
        for run in p0.runs[1:]:
            run.text = ""
    else:
        p0.text = text
    for p in cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)


def set_para(p: Paragraph, text: str) -> None:
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.text = text


def find_para(doc: Document, contains: str = None, startswith: str = None):
    for p in doc.paragraphs:
        t = p.text or ""
        if startswith is not None and t.startswith(startswith):
            return p
        if contains is not None and contains in t:
            return p
    raise KeyError(contains or startswith)


def insert_paragraph_after(paragraph: Paragraph, text: str) -> Paragraph:
    new_p = deepcopy(paragraph._p)
    for child in list(new_p):
        if child.tag == qn("w:r"):
            new_p.remove(child)
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if paragraph.runs:
        run = new_para.add_run(text)
        src = paragraph.runs[0]
        run.bold = src.bold
        run.italic = src.italic
        if src.font.size:
            run.font.size = src.font.size
        if src.font.name:
            run.font.name = src.font.name
    else:
        new_para.add_run(text)
    return new_para


def ensure_rows(table: Table, n: int) -> None:
    while len(table.rows) < n:
        table._tbl.append(deepcopy(table.rows[-1]._tr))
    while len(table.rows) > n:
        table.rows[-1]._tr.getparent().remove(table.rows[-1]._tr)


def fill_table(table: Table, rows: list[list[str]]) -> None:
    ensure_rows(table, len(rows))
    for i, vals in enumerate(rows):
        for cell, val in zip(table.rows[i].cells, vals):
            set_cell_text(cell, val)


def main() -> None:
    doc = Document(str(DOCX))

    if any("Tabla 7d." in (p.text or "") for p in doc.paragraphs):
        print("CAPEX/OPEX ya insertado; no se duplica.")
        return

    # Ancla: después del párrafo de medición o de Tabla 7c
    try:
        anchor = find_para(doc, startswith="Tabla 7c.")
    except KeyError:
        anchor = find_para(doc, contains="Escenario C — Redimensionar")

    # Si hay párrafo de medición después de 7c, insertar después de ese
    try:
        med = find_para(doc, contains="Medición (válida en A, B o C)")
        anchor = med
    except KeyError:
        pass

    p = insert_paragraph_after(
        anchor,
        "4.1.2. Dimensión CAPEX y OPEX (lectura para dirección y contraparte)",
    )
    p = insert_paragraph_after(
        p,
        "Para que el cliente (dirección / contraparte) vea el cambio en la misma lógica "
        "financiera que suele usar un piloto, los costos de C-05b y C-06 se presentan también "
        "como CAPEX y OPEX. CAPEX, aquí, es el esfuerzo e inversión de puesta en marcha que se "
        "concentra en la ventana del proyecto (diseño UI/UX, integración Gemini, despliegue "
        "inicial en Google Cloud, corpus FAQ). OPEX es el gasto recurrente de operación mientras "
        "el chatbot y el tablero estén en uso (tokens Gemini, Cloud Run u hosting, logs, y las "
        "horas de soporte humano que siguen siendo necesarias en zona compleja). En software "
        "cloud casi todo es OPEX contable; igual se separa la «puesta en marcha del cambio» de "
        "la «cuenta que sigue mes a mes», para no vender el chatbot como un gasto único ni "
        "esconder la factura de API.",
    )

    # Plantilla: tabla con Gemini (7b)
    template = None
    for table in doc.tables:
        blob = " ".join(c.text for r in table.rows for c in r.cells)
        if "Gemini" in blob and "Rubro" in blob:
            template = table
            break
    if template is None:
        template = doc.tables[2]

    new_t = Table(deepcopy(template._tbl), p._parent)
    p._p.addnext(new_t._tbl)

    fill_table(
        new_t,
        [
            [
                "Concepto",
                "CAPEX / OPEX",
                "US$ en ventana piloto (orden)",
                "Qué ve el cliente",
                "Quién controla",
            ],
            [
                "Diseño UI/UX cola (C-06) + chat FAQ",
                "CAPEX (puesta en marcha)",
                "0 caja (HH equipo: ~6–10 pd)",
                "Inversión de construcción del cambio; no es una cuota mensual aparte",
                "Giannina / Nicolás",
            ],
            [
                "Integración Gemini + despliegue inicial en Google Cloud",
                "CAPEX (puesta en marcha)",
                "0 caja HH (~5–8 pd) + setup GCloud (~US$ 40–80 si hay cuotas de alta)",
                "Dejar el módulo listo en el ambiente del piloto",
                "Nicolás / Ari",
            ],
            [
                "Corpus FAQ + pruebas de calidad",
                "CAPEX (puesta en marcha)",
                "0 caja (HH: ~2–4 pd)",
                "Contenido curado antes de abrir el chat a usuarios",
                "Skarlett",
            ],
            [
                "API Gemini (tokens) en nov–dic",
                "OPEX (recurrente)",
                "60–120 en la ventana; luego ~US$ 20–80/mes según uso",
                "Cuenta que escala con las conversaciones; tope en Excel",
                "Ari",
            ],
            [
                "Google Cloud Run / hosting / logs del chat",
                "OPEX (recurrente)",
                "60–100 en la ventana; luego margen mensual bajo en piloto",
                "Infra que sigue mientras el servicio esté arriba",
                "Ari",
            ],
            [
                "Soporte humano local (sigue existiendo)",
                "OPEX (recurrente)",
                "Parte del variable ~760 del plan; no se elimina",
                "El chatbot busca bajarlo; la zona compleja exige mantener capacidad",
                "Giannina / soporte",
            ],
            [
                "Colchón sobrecosto IA (Escenario C)",
                "OPEX / reserva",
                "80–120 dentro del +500 del techo 21.200",
                "Amortigua picos de uso sin inventar un segundo techo informal",
                "Ari / Giannina",
            ],
        ],
    )

    # Caption
    cap_el = deepcopy(p._p)
    for child in list(cap_el):
        if child.tag == qn("w:r"):
            cap_el.remove(child)
    new_t._tbl.addnext(cap_el)
    cap = Paragraph(cap_el, p._parent)
    run = cap.add_run(
        "Tabla 7d. Lectura CAPEX / OPEX de C-05b y C-06 para dirección y contraparte. "
        "CAPEX = puesta en marcha del cambio en la ventana; OPEX = operación recurrente "
        "(Gemini + Google Cloud + soporte humano)."
    )
    if p.runs and p.runs[0].font.size:
        run.font.size = p.runs[0].font.size
    if p.runs and p.runs[0].font.name:
        run.font.name = p.runs[0].font.name

    p2 = insert_paragraph_after(
        cap,
        "Mensaje al cliente: aprobar C-05b en Escenario C (techo US$ 21.200) autoriza un "
        "CAPEX de puesta en marcha (trabajo del equipo ya en plan, más setup GCloud) y un "
        "OPEX acotado de Gemini y Google Cloud, medido cada dos semanas. El beneficio "
        "esperado es reducir OPEX de soporte informal; eso se comprueba con la línea base "
        "post 9 de noviembre, no se promete como ahorro fijo en el acta. Las fases "
        "posteriores (monitoreo, alertas, IA de flota) serían OPEX/CAPEX nuevos y exigen "
        "otro control de cambios.",
    )

    # Tabla 11 — reforzar lectura cliente
    t11 = doc.tables[-1]
    if "Pregunta" in (t11.rows[0].cells[0].text or ""):
        # Añadir fila si no existe pregunta CAPEX
        has = any("CAPEX" in (r.cells[0].text or "") for r in t11.rows)
        if not has:
            tr = deepcopy(t11.rows[-1]._tr)
            t11._tbl.append(tr)
            set_cell_text(
                t11.rows[-1].cells[0],
                "¿Cómo se lee el costo para el cliente (CAPEX/OPEX)?",
            )
            set_cell_text(
                t11.rows[-1].cells[1],
                "CAPEX: puesta en marcha de C-06 y C-05b (UI/UX, integración Gemini, "
                "despliegue GCloud, FAQ) con HH del equipo. OPEX: Gemini + Google Cloud "
                "recurrentes y soporte humano que se mantiene. Escenario C (techo 21.200) "
                "financia ese OPEX acotado sin vaciar soporte. Detalle en Tabla 7d.",
            )

    doc.save(str(DOCX))
    print("OK", DOCX)


if __name__ == "__main__":
    main()
